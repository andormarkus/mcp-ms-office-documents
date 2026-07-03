"""Tests for the create-side DOCX gap features (page setup, fonts, columns,
tab stops, bookmarks, footnotes, table shading, image alt text, smart quotes).

These implement, in pure Python, the create-side ideas documented in the
skills/docx SKILL.md (page size/orientation, default fonts, internal links,
tab stops, multi-column layouts, table shading, image alt text, smart quotes)
without depending on LibreOffice or docx-js.
"""
import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

import pytest
from docx import Document
from docx.oxml.ns import qn

from docx_tools.base_docx_tool import _markdown_to_doc
from docx_tools.directives import parse_document_directives
from docx_tools.patterns import apply_smart_quotes

OUTPUT_DIR = Path(__file__).parent / "output" / "docx"


@pytest.fixture(scope="module", autouse=True)
def setup_output_dir():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    yield


def _doc(markdown, **kwargs):
    return _markdown_to_doc(markdown, **kwargs)


# =============================================================================
# Directive parsing
# =============================================================================

class TestParseDocumentDirectives:
    def test_no_directives(self):
        directives, body = parse_document_directives("# Title\n\nBody text.")
        assert directives == {}
        assert body == "# Title\n\nBody text."

    def test_single_directive_removed(self):
        content = "<!-- page: letter -->\n# Title\n\nBody."
        directives, body = parse_document_directives(content)
        assert directives == {"page": "letter"}
        assert "<!-- page:" not in body
        assert "# Title" in body

    def test_multiple_directives(self):
        content = (
            "<!-- page: letter -->\n"
            "<!-- orientation: landscape -->\n"
            "<!-- font: Arial -->\n"
            "# Title\n"
        )
        directives, body = parse_document_directives(content)
        assert directives == {"page": "letter", "orientation": "landscape", "font": "Arial"}
        assert "# Title" in body

    def test_non_document_level_directive_left_untouched(self):
        """A block directive (e.g. borderless) in the leading run is preserved
        so the per-block mechanism still sees it."""
        content = "<!-- borderless -->\n| A | B |\n|---|---|\n| 1 | 2 |\n"
        directives, body = parse_document_directives(content)
        assert directives == {}
        assert "<!-- borderless -->" in body

    def test_mixed_document_and_block_directives(self):
        content = (
            "<!-- page: letter -->\n"
            "<!-- borderless -->\n"
            "| A | B |\n|---|---|\n| 1 | 2 |\n"
        )
        directives, body = parse_document_directives(content)
        assert directives == {"page": "letter"}
        assert "<!-- borderless -->" in body
        assert "<!-- page:" not in body

    def test_stops_at_first_content_line(self):
        """Directives after real content are not treated as document-level."""
        content = "# Title\n\n<!-- page: letter -->\n\nMore text."
        directives, body = parse_document_directives(content)
        assert directives == {}
        assert body == content

    def test_empty_content(self):
        directives, body = parse_document_directives("")
        assert directives == {}
        assert body == ""


# =============================================================================
# Page size / orientation / margins
# =============================================================================

class TestPageSetup:
    def test_page_size_letter_directive(self):
        doc = _doc("<!-- page: letter -->\n\n# Test")
        section = doc.sections[0]
        assert int(section.page_width) == int(Emu_from_twips(12240))
        assert int(section.page_height) == int(Emu_from_twips(15840))

    def test_no_page_directive_leaves_template_default_untouched(self):
        """With no 'page'/'orientation'/'margin' directive or parameter,
        apply_page_setup must be a no-op — page size stays whatever the
        (possibly custom) template already defines."""
        from docx import Document as _Document
        from docx_tools.document_features import load_templates

        template_path = load_templates()
        baseline = _Document(template_path) if template_path else _Document()
        doc = _doc("# Test")
        assert int(doc.sections[0].page_width) == int(baseline.sections[0].page_width)
        assert int(doc.sections[0].page_height) == int(baseline.sections[0].page_height)

    def test_page_size_legal(self):
        doc = _doc("<!-- page: legal -->\n\n# Test")
        section = doc.sections[0]
        assert int(section.page_width) == int(Emu_from_twips(12240))
        assert int(section.page_height) == int(Emu_from_twips(20160))

    def test_orientation_landscape_swaps_dimensions(self):
        doc = _doc("<!-- page: letter -->\n<!-- orientation: landscape -->\n\n# Test")
        section = doc.sections[0]
        assert section.page_width > section.page_height
        assert int(section.page_width) == int(Emu_from_twips(15840))
        assert int(section.page_height) == int(Emu_from_twips(12240))

    def test_page_size_param_overrides_directive(self):
        doc = _doc("<!-- page: a4 -->\n\n# Test", page_size="letter")
        section = doc.sections[0]
        assert int(section.page_width) == int(Emu_from_twips(12240))

    def test_orientation_param_overrides_directive(self):
        doc = _doc("<!-- orientation: portrait -->\n\n# Test", orientation="landscape")
        section = doc.sections[0]
        assert section.page_width > section.page_height

    def test_margin_all_sides(self):
        from docx.shared import Inches
        doc = _doc("<!-- margin: 2in -->\n\n# Test")
        section = doc.sections[0]
        assert section.top_margin == Inches(2)
        assert section.left_margin == Inches(2)

    def test_margin_per_side(self):
        from docx.shared import Inches
        doc = _doc("<!-- margin: top=1in bottom=1.5in left=1.25in right=1.25in -->\n\n# Test")
        section = doc.sections[0]
        assert section.top_margin == Inches(1)
        assert section.bottom_margin == Inches(1.5)
        assert section.left_margin == Inches(1.25)

    def test_no_directive_leaves_default(self):
        doc = _doc("# Test")
        assert doc.sections[0].orientation is not None  # sanity: doesn't crash


def Emu_from_twips(twips):
    from docx.shared import Twips
    return Twips(twips)


# =============================================================================
# Default font / heading font
# =============================================================================

class TestDefaultFont:
    def test_font_directive_applied_to_normal_style(self):
        doc = _doc("<!-- font: Arial -->\n\nSome text.")
        assert doc.styles['Normal'].font.name == 'Arial'

    def test_font_param_overrides_directive(self):
        doc = _doc("<!-- font: Arial -->\n\nSome text.", default_font="Georgia")
        assert doc.styles['Normal'].font.name == 'Georgia'

    def test_font_size_directive(self):
        from docx.shared import Pt
        doc = _doc("<!-- font-size: 14 -->\n\nSome text.")
        assert doc.styles['Normal'].font.size == Pt(14)

    def test_heading_font_directive_applies_to_all_heading_levels(self):
        doc = _doc("<!-- heading-font: Georgia -->\n\n# H1\n\n## H2")
        assert doc.styles['Heading 1'].font.name == 'Georgia'
        assert doc.styles['Heading 2'].font.name == 'Georgia'

    def test_font_directive_also_applies_to_headings_by_default(self):
        doc = _doc("<!-- font: Arial -->\n\n# H1")
        assert doc.styles['Heading 1'].font.name == 'Arial'

    def test_no_font_directive_leaves_default(self):
        doc = _doc("# Test")
        # Should not raise; Normal style font name may be None (inherits template default)
        assert doc.styles['Normal'] is not None


# =============================================================================
# Multi-column layout
# =============================================================================

class TestColumns:
    def test_columns_two(self):
        doc = _doc("<!-- columns: 2 -->\n\nSome text.")
        sectPr = doc.sections[-1]._sectPr
        cols = sectPr.find(qn('w:cols'))
        assert cols is not None
        assert cols.get(qn('w:num')) == '2'

    def test_columns_with_separator(self):
        doc = _doc("<!-- columns: 2 sep -->\n\nSome text.")
        sectPr = doc.sections[-1]._sectPr
        cols = sectPr.find(qn('w:cols'))
        assert cols.get(qn('w:sep')) == '1'

    def test_columns_custom_widths(self):
        doc = _doc("<!-- columns: custom:5400,3240 -->\n\nSome text.")
        sectPr = doc.sections[-1]._sectPr
        cols = sectPr.find(qn('w:cols'))
        assert cols.get(qn('w:equalWidth')) == '0'
        col_elements = cols.findall(qn('w:col'))
        assert len(col_elements) == 2
        assert col_elements[0].get(qn('w:w')) == '5400'

    def test_no_columns_directive_leaves_template_default_untouched(self):
        """With no 'columns' directive, apply_columns must be a no-op — any
        <w:cols> already present in the (possibly custom) template is left
        exactly as-is."""
        from docx import Document as _Document
        from docx_tools.document_features import load_templates

        template_path = load_templates()
        baseline = _Document(template_path) if template_path else _Document()
        baseline_cols = baseline.sections[-1]._sectPr.find(qn('w:cols'))

        doc = _doc("# Test")
        cols = doc.sections[-1]._sectPr.find(qn('w:cols'))
        if baseline_cols is None:
            assert cols is None
        else:
            assert cols is not None
            assert cols.get(qn('w:num')) == baseline_cols.get(qn('w:num'))


# =============================================================================
# Footnotes
# =============================================================================

class TestFootnotes:
    def test_footnote_reference_emitted(self):
        doc = _doc("This has a footnote[^This is the footnote text.] in it.")
        xml = doc.element.xml
        assert 'w:footnoteReference' in xml

    def test_footnotes_part_created(self):
        from lxml import etree
        doc = _doc("Revenue grew[^Source: Annual Report 2026.] significantly.")
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        part = doc.part.part_related_by(RT.FOOTNOTES)
        assert part is not None
        xml = etree.tostring(part.element).decode('utf-8')
        assert 'Source: Annual Report 2026.' in xml

    def test_multiple_footnotes_get_incrementing_ids(self):
        doc = _doc("First[^one] and second[^two] footnotes.")
        xml = doc.element.xml
        assert 'w:id="1"' in xml
        assert 'w:id="2"' in xml

    def test_footnote_does_not_break_surrounding_text(self):
        doc = _doc("Before footnote[^note text] after footnote.")
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Before footnote" in full_text
        assert "after footnote" in full_text


# =============================================================================
# Internal hyperlinks / bookmarks
# =============================================================================

class TestInternalLinksAndBookmarks:
    def test_internal_hyperlink_uses_anchor_not_relationship(self):
        doc = Document()
        p = doc.add_paragraph()
        from docx_tools.inline_formatting import parse_inline_formatting
        parse_inline_formatting("[See Chapter 1](#chapter1)", p)
        hyperlink = p._p.find(qn('w:hyperlink'))
        assert hyperlink is not None
        assert hyperlink.get(qn('w:anchor')) == 'chapter1'
        assert hyperlink.get(qn('r:id')) is None

    def test_external_hyperlink_still_uses_relationship(self):
        doc = Document()
        p = doc.add_paragraph()
        from docx_tools.inline_formatting import parse_inline_formatting
        parse_inline_formatting("[Visit](https://example.com)", p)
        hyperlink = p._p.find(qn('w:hyperlink'))
        assert hyperlink is not None
        assert hyperlink.get(qn('w:anchor')) is None
        assert hyperlink.get(qn('r:id')) is not None

    def test_bookmark_directive_creates_bookmark_pair(self):
        doc = _doc("<!-- bookmark: chapter1 -->\n# Chapter 1\n\nContent.")
        xml = doc.element.xml
        assert 'w:bookmarkStart' in xml
        assert 'w:bookmarkEnd' in xml
        assert 'w:name="chapter1"' in xml

    def test_bookmark_and_internal_link_roundtrip(self):
        markdown = (
            "<!-- bookmark: chapter1 -->\n"
            "# Chapter 1\n\n"
            "See [Chapter 1](#chapter1) for details.\n"
        )
        doc = _doc(markdown)
        xml = doc.element.xml
        assert 'w:name="chapter1"' in xml
        assert 'w:anchor="chapter1"' in xml


# =============================================================================
# Tab stops
# =============================================================================

class TestTabStops:
    def test_tab_character_creates_tab_run(self):
        doc = Document()
        p = doc.add_paragraph()
        from docx_tools.inline_formatting import parse_inline_formatting
        parse_inline_formatting("Company Name\tJanuary 2026", p)
        xml = p._p.xml
        assert '<w:tab/>' in xml

    def test_tab_directive_right_with_dot_leader(self):
        doc = _doc("<!-- tab: right dot -->\n\nIntroduction\tp. 3")
        # Find the paragraph containing the tab stop definition.
        found = False
        for p in doc.paragraphs:
            pPr = p._p.find(qn('w:pPr'))
            if pPr is None:
                continue
            tabs = pPr.find(qn('w:tabs'))
            if tabs is not None:
                tab = tabs.find(qn('w:tab'))
                assert tab.get(qn('w:val')) == 'right'
                assert tab.get(qn('w:leader')) == 'dot'
                found = True
        assert found, "Expected a paragraph with an explicit right/dot tab stop"

    def test_tab_directive_default_alignment(self):
        doc = _doc("<!-- tab: right -->\n\nLabel\tValue")
        xml = doc.element.xml
        assert 'w:tabs' in xml
        assert 'w:val="right"' in xml


# =============================================================================
# Table shading
# =============================================================================

class TestTableShading:
    def test_header_shading(self):
        markdown = (
            "<!-- shade: header=D5E8F0 -->\n"
            "| Name | Age |\n|------|-----|\n| John | 25 |\n"
        )
        doc = _doc(markdown)
        table = doc.tables[0]
        header_tcPr = table.cell(0, 0)._tc.find(qn('w:tcPr'))
        shd = header_tcPr.find(qn('w:shd'))
        assert shd is not None
        assert shd.get(qn('w:fill')) == 'D5E8F0'
        assert shd.get(qn('w:val')) == 'clear'
        # Data row should be unshaded.
        data_tcPr = table.cell(1, 0)._tc.find(qn('w:tcPr'))
        if data_tcPr is not None:
            assert data_tcPr.find(qn('w:shd')) is None

    def test_alt_row_shading(self):
        markdown = (
            "<!-- shade: alt=F2F2F2 -->\n"
            "| Name | Age |\n|------|-----|\n"
            "| John | 25 |\n| Jane | 30 |\n| Bob | 35 |\n"
        )
        doc = _doc(markdown)
        table = doc.tables[0]
        # Row 1 (first data row) shaded, row 2 not, row 3 shaded (0-indexed rows 1,2,3).
        row1_tcPr = table.cell(1, 0)._tc.find(qn('w:tcPr'))
        assert row1_tcPr is not None and row1_tcPr.find(qn('w:shd')) is not None
        row3_tcPr = table.cell(3, 0)._tc.find(qn('w:tcPr'))
        assert row3_tcPr is not None and row3_tcPr.find(qn('w:shd')) is not None

    def test_header_and_alt_combined(self):
        markdown = (
            "<!-- shade: header=D5E8F0 alt=F2F2F2 -->\n"
            "| Name | Age |\n|------|-----|\n| John | 25 |\n| Jane | 30 |\n"
        )
        doc = _doc(markdown)
        table = doc.tables[0]
        header_shd = table.cell(0, 0)._tc.find(qn('w:tcPr')).find(qn('w:shd'))
        assert header_shd.get(qn('w:fill')) == 'D5E8F0'

    def test_invalid_color_ignored(self):
        markdown = (
            "<!-- shade: header=notacolor -->\n"
            "| Name | Age |\n|------|-----|\n| John | 25 |\n"
        )
        doc = _doc(markdown)  # Should not raise.
        assert doc is not None


# =============================================================================
# Image alt text
# =============================================================================

class TestImageAltText:
    def test_image_alt_text_set_on_docpr(self):
        markdown = "![A descriptive caption](https://invalid-test-domain.test/img.png)"
        doc = _doc(markdown)
        xml = doc.element.xml
        # Falls back to error placeholder text since the URL is invalid, but
        # verify the alt-text plumbing itself via a direct unit test below.
        assert "Image could not be loaded" in xml or "descriptive caption" in xml

    def test_set_image_alt_text_helper_sets_docpr_attrs(self):
        import base64
        import io as _io
        from docx.shared import Inches
        from docx_tools.block_elements import _set_image_alt_text

        doc = Document()
        run = doc.add_paragraph().add_run()
        png_bytes = base64.b64decode(
            'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY'
            '42YAAAAASUVORK5CYII='
        )
        picture = run.add_picture(_io.BytesIO(png_bytes), width=Inches(1))
        _set_image_alt_text(picture, "A descriptive caption")
        docPr = picture._inline.find(qn('wp:docPr'))
        assert docPr.get('descr') == "A descriptive caption"
        assert docPr.get('title') == "A descriptive caption"


# =============================================================================
# Smart quotes
# =============================================================================

class TestSmartQuotes:
    def test_apply_smart_quotes_basic(self):
        result = apply_smart_quotes('She said "hello" to me.')
        assert result == 'She said \u201chello\u201d to me.'

    def test_apply_smart_quotes_apostrophe(self):
        result = apply_smart_quotes("It's a test.")
        assert result == "It\u2019s a test."

    def test_apply_smart_quotes_skips_inline_code(self):
        result = apply_smart_quotes('Use `it\'s "code"` literally.')
        assert '`it\'s "code"`' in result

    def test_apply_smart_quotes_skips_fenced_code_blocks(self):
        content = '```\nvar x = "hello";\n```\n'
        result = apply_smart_quotes(content)
        assert '"hello"' in result

    def test_smart_quotes_off_by_default(self):
        doc = _doc('She said "hello" to me.')
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert '"hello"' in full_text

    def test_smart_quotes_directive_on(self):
        doc = _doc('<!-- smart-quotes: on -->\n\nShe said "hello" to me.')
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "\u201chello\u201d" in full_text

    def test_smart_quotes_param_overrides_directive(self):
        doc = _doc('<!-- smart-quotes: off -->\n\nShe said "hello" to me.', smart_quotes=True)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "\u201chello\u201d" in full_text


# =============================================================================
# Backward compatibility (no directives => unchanged rendering)
# =============================================================================

class TestBackwardCompatibility:
    def test_document_without_directives_renders_normally(self):
        markdown = """# Report

## Summary

This is a **test** document with a [link](https://example.com).

| A | B |
|---|---|
| 1 | 2 |
"""
        doc = _doc(markdown)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Report" in full_text
        assert "Summary" in full_text
        assert len(doc.tables) == 1

    def test_existing_borderless_widths_directives_still_work(self):
        markdown = (
            "<!-- borderless -->\n"
            "<!-- widths: 30 70 -->\n"
            "| A | B |\n|---|---|\n| 1 | 2 |\n"
        )
        doc = _doc(markdown)
        table = doc.tables[0]
        tblPr = table._tbl.tblPr
        borders = tblPr.find(qn('w:tblBorders'))
        assert borders is not None


if __name__ == "__main__":
    pytest.main([__file__, "-v", "--tb=short"])
